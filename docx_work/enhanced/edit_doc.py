
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src=Path(r'C:\Users\wangmeng\Desktop\智巡云_Web端作品功能与原型设计.docx')
out=Path(r'C:\Users\wangmeng\Desktop\智巡云_Web端作品功能与原型设计_精进版.docx')
fig=Path(r'C:\Users\wangmeng\Desktop\git\docx_work\enhanced\figures')

def font(run,size=12,bold=None,color='000000'):
 n='仿宋_GB2312'
 run.font.name=n
 rpr=run._element.get_or_add_rPr()
 rf=rpr.rFonts
 if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
 for k in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+k),n)
 run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
 if bold is not None: run.bold=bold

def set_text(p,text):
 p.clear(); p.add_run(text)
def add_para(doc,text,boldlead=None):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 pf=p.paragraph_format; pf.space_before=Pt(0); pf.space_after=Pt(7); pf.line_spacing=1.35; pf.widow_control=True
 if boldlead and text.startswith(boldlead):
  a=p.add_run(boldlead); font(a,bold=True); b=p.add_run(text[len(boldlead):]); font(b)
 else:
  a=p.add_run(text); font(a)
 return p
def add_head(doc,text):
 p=doc.add_paragraph(style='Heading 1'); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 pf=p.paragraph_format; pf.space_before=Pt(13); pf.space_after=Pt(7); pf.line_spacing=1.2; pf.keep_with_next=True
 r=p.add_run(text); font(r,bold=True)
 return p
def margins(cell,top=100,start=125,bottom=100,end=125):
 pr=cell._tc.get_or_add_tcPr(); mar=pr.find(qn('w:tcMar'))
 if mar is None: mar=OxmlElement('w:tcMar'); pr.append(mar)
 for key,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  e=mar.find(qn('w:'+key))
  if e is None:e=OxmlElement('w:'+key);mar.append(e)
  e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa')
def shade(cell,fill):
 pr=cell._tc.get_or_add_tcPr(); e=pr.find(qn('w:shd'))
 if e is None:e=OxmlElement('w:shd');pr.append(e)
 e.set(qn('w:fill'),fill)
def add_fig(doc,path,caption,width=6.0):
 t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Inches(6.14)
 c=t.cell(0,0); margins(c,60,60,60,60); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4); p.paragraph_format.keep_with_next=True
 p.add_run().add_picture(str(path),width=Inches(width))
 cp=c.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after=Pt(6); cp.paragraph_format.keep_together=True
 r=cp.add_run(caption); font(r,bold=True)
 return t
def add_grid(doc,headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,x in enumerate(headers):
  c=t.rows[0].cells[i];c.text='';shade(c,'DCE6F1');margins(c)
  p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  r=p.add_run(x);font(r,bold=True)
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):
   c=cells[i];c.text='';margins(c);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(2)
   r=p.add_run(x);font(r,bold=(i==0))
 return t

d=Document(str(src))
repl={
1:'“智巡云——基于边缘感知的路网健康巡检系统”Web 管理后台面向交通管理部门的路网巡检、告警审核、工单调度与车载设备运维场景。本文在现有页面原型基础上，进一步说明产品目标、系统边界、主要角色、功能结构、页面操作、数据流转和异常处理方式。文中界面截图来自 Web 原型；新增示意图用于解释系统关系与业务机制。',
12:'系统围绕管理人员的日常工作顺序组织页面：登录认证后进入路网总览，随后可查看事件中心、地图巡检、设备中心和工单详情；处理材料上传及智巡助手作为工单与分析过程中的辅助能力。以下从页面职责、主要信息和操作反馈三个方面说明原型设计。',
13:'登录页负责建立业务入口。用户输入账号和密码后，前端调用登录接口；成功后保存令牌及必要的用户资料，再加载业务应用。账号格式错误、密码错误或接口超时均在表单附近显示反馈；用户修正后可重新提交。退出登录会清除会话并回到登录页。',
14:'路网总览用于快速判断当前运行态势。页面将工单状态、事件分布、设备健康和趋势信息组合展示；地图为事件提供空间上下文，分析区帮助识别告警与设备的变化。数据通过 api.js 并行读取，图表和指标采用原位更新；页面切换后旧请求结果会被丢弃，避免过时响应覆盖当前状态。',
15:'事件中心采用“处置队列 + 工单详情”的双栏结构。队列支持关键词、状态、优先级、日期和排序筛选，状态数量与列表同步变化；详情区集中展示事件描述、地点、上报时间、相关材料、处理进度和可用操作。审核、派发或催办成功后，相关数据缓存失效并重新读取，页面同步更新状态。',
16:'地图巡检页以腾讯地图为主视图，展示有有效经纬度的事件和人员位置。告警点按优先级区分样式，近距离点位通过聚合减少遮挡；用户可按事件类别开关图层，点击单点进入对应详情，点击聚合点则调整地图视野。无坐标记录不参与地图绘制，但仍保留在列表和工单详情中。',
17:'设备中心从设备总量、状态分布和单设备明细三个层次展示运维信息。设备按照正常、异常、离线状态归类，状态图例可打开对应设备清单；详情内容包含设备编号、车辆名称、温度、存储和后端提供的异常说明。未知状态单独提示，不强行计入已识别状态比例。',
18:'处理材料区域只对允许上传的工单状态开放。用户可批量选择 JPG、PNG、WEBP 图片和 MP4 视频，每批最多 10 个文件、单个文件不超过 200 MB；PNG/WEBP 在浏览器端转换为白底 JPG，再由浏览器直传 COS。队列逐项显示进度和错误信息，失败项可以重试；上传成功先显示临时预览，既有扫描与查询接口随后完成持久同步。',
19:'个人信息入口展示当前登录资料并提供退出操作。会话变化时，读取缓存按 API 地址和令牌作用域清空；页面控制器在离开页面时销毁事件监听与定时器。总览页和设备页分别按页面策略刷新，避免后台页面在不可见或已离开后继续无效请求。'
}
for i,s in repl.items():set_text(d.paragraphs[i],s)
# Expanded feature descriptions in the existing two-column function matrix.
features={
'登录认证':'用户使用账号和密码登录。前端调用 /api/login，成功后保存服务端令牌和必要的用户资料；业务请求统一携带 Bearer 令牌。登录失败、超时或会话被拒绝时显示反馈并清理失效状态；主动退出时清除会话后返回登录页。',
'路网总览':'汇总今日告警、待处理工单、设备在线情况、工单状态分布、趋势和地图事件。页面并行读取所需数据；指标、图表和地图标记原位更新，按当前页面策略定时刷新。导航切换或数据版本变化时，丢弃过期异步响应。',
'事件中心':'以处置队列与详情双栏工作区呈现事件。支持关键词、状态、优先级、日期和排序筛选；列表数量、状态计数和详情选择保持同步。审核、发送工单、催办等操作写回后端，成功后刷新相关数据。',
'地图巡检':'通过腾讯地图 GL JS 展示有坐标的告警、聚合点和人员位置。点位按优先级显示样式，分类图层可筛选；点击聚合点放大视野，点击单点打开事件详情。无效坐标、地图 Key 缺失或 SDK 加载失败时显示降级提示。',
'设备中心':'按正常、异常、离线状态统计车辆设备，并提供状态明细、设备列表和异常原因。设备卡片展示编号、车辆、温度、存储和后端状态；未知状态单独提示。点击设备可进一步查看运行详情。',
'处理材料':'工单详情支持多选 JPG、PNG、WEBP 图片及 MP4 视频。前端验证格式、MIME、文件数量和大小，PNG/WEBP 转为 JPG 后由浏览器直传腾讯云 COS。队列逐项展示进度；失败文件保留并支持重试，上传成功后显示临时预览。',
'智巡助手':'通过 Dify 发布的嵌入页面提供智能问答入口。用户主动打开时才创建 iframe，支持关闭、重试、独立窗口和面板拖动/尺寸调整。跨域环境不读取 iframe 内部业务状态，连接提示依据加载、超时或错误事件。',
'个人信息':'从当前登录会话显示姓名、部门等资料，提供退出登录操作。退出后清除会话并重新进入登录流程；个人资料用于展示，不代替后端权限校验。',
'发布与部署':'本地可使用静态服务或项目 Node 服务预览，构建脚本将 HTML、CSS、JavaScript 与资源复制到 web-dist。生产发布需配置 API 地址、腾讯地图域名白名单、COS CORS 和 Dify 嵌入代理。'
}
for row in d.tables[0].rows[1:]:
 key=row.cells[0].text.strip()
 if key in features: row.cells[1].text=features[key]
# Fit the long feature matrix to the available text width and repeat semantic table headers.
feature_table=d.tables[0]
feature_table.autofit=False
for i,w in enumerate([1.55,4.59]):
 feature_table.columns[i].width=Inches(w)
 for row in feature_table.rows:
  row.cells[i].width=Inches(w)
for table_index in [0]:
 row= d.tables[table_index].rows[0]
 trPr=row._tr.get_or_add_trPr()
 if not trPr.xpath('./w:tblHeader'): trPr.append(OxmlElement('w:tblHeader'))
# Use semantic Word styles while keeping the requested 12 pt typography.
d.paragraphs[0].style=d.styles['Title']
d.paragraphs[2].style=d.styles['Heading 1']
d.paragraphs[4].style=d.styles['Heading 1']
d.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
for p in d.paragraphs:
 if p.text.strip() in ['功能概述','原型设计']:
  p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_with_next=True
  for r in p.runs:r.bold=True
# Technical sections and illustrations
add_head(d,'一、总体设计思路')
add_para(d,'系统以“事件数据为主线、空间位置为上下文、工单状态为闭环、设备健康为保障”为设计思路。路网总览承担全局态势入口，事件中心处理审核与派单，地图巡检提供空间查看，设备中心负责车辆设备运维；材料上传与智能助手分别补充现场凭证和自然语言查询能力。')
add_para(d,'在实现边界上，浏览器负责页面呈现和用户交互，api.js 集中处理业务 API 与数据模型转换；腾讯地图、COS 和 Dify 通过独立接入点连接。业务状态最终以服务端响应为准，前端缓存只用于短时复用和减少重复读取。')
add_fig(d,fig/'architecture.png','图 8 Web 端总体架构与外部服务关系')
add_head(d,'二、核心业务流程')
add_para(d,'从告警到闭环的典型流程由事件接入、人工审核、工单派发、现场处置和结果回查组成。队列页面帮助用户识别优先级；工单详情汇集事件信息和证据；处理材料上传后，既有对象扫描流程负责将材料关联回工单。具体审核条件与派单规则以业务接口和后端配置为准。')
add_fig(d,fig/'closed-loop.png','图 9 告警审核与工单处置闭环')
add_para(d,'页面反馈采用“发起操作、等待确认、成功后刷新、失败时保留上下文”的方式。上传等耗时操作显示逐项进度，并在运行期间阻止误切换；网络错误时保留队列和筛选条件，供用户重试。')
add_head(d,'三、关键页面与交互状态')
add_para(d,'路网总览提供态势摘要，不承担复杂编辑。用户从指标或事件列表进入相应业务页面后，系统显示当前任务需要的上下文；返回总览时再以最新读取结果更新指标、图表和状态。')
add_para(d,'事件中心需覆盖空列表、无筛选结果、接口失败、操作处理中、操作成功和操作失败等状态。详情区与队列选中项保持一致；高优先级和待处置状态使用文字与颜色共同提示，避免只依赖颜色传达信息。')
add_para(d,'地图巡检需覆盖加载中、正常显示、无有效坐标、Key 未配置和网络异常等情形。图层筛选只改变地图上可见类别，人员图层与告警聚合分开维护；事件列表仍可作为无坐标记录的查看入口。')
add_para(d,'设备中心分别统计正常、异常、离线和未知状态。状态明细展示设备名称、编号与异常原因；刷新数据时保留当前选中状态和滚动位置，减少重复定位。')
add_head(d,'四、数据刷新与一致性控制')
add_para(d,'read-cache.js 在登录会话和 API 地址作用域内短时缓存成功结果，并复用进行中的同类请求；写操作完成后 api.js 使相关资源失效。页面控制器在异步响应返回时检查导航编号和数据版本，避免用户已切页或完成新操作后，较早响应覆盖当前界面。总览和设备页按各自策略周期读取，离开页面后停止对应刷新任务。')
add_fig(d,fig/'refresh.png','图 10 短缓存、写后失效与异步响应校验')
add_head(d,'五、工单材料上传设计')
add_para(d,'浏览器端上传链路检查登录令牌、文件类型、空文件、批次数量和单文件大小；PNG/WEBP 使用浏览器图像能力转换为白底 JPG，视频按 MP4 处理。文件通过 XMLHttpRequest 上传并反馈逐项进度；对象键包含工单编号和 UUID，避免重名并兼容既有材料扫描任务。')
add_para(d,'上传成功后页面立即显示临时预览，之后由现有扫描与查询接口完成持久同步。传输失败时显示错误信息并保留该项供重试；上传过程中不允许移除队列项目，减少重复提交和上下文丢失。')
add_fig(d,fig/'cos-upload.png','图 11 浏览器直传 COS 的处理链路')
add_head(d,'六、异常处理与降级策略')
rows=[
('登录与会话','接口失败或超时：保留输入并显示错误；令牌被拒绝：清理失效会话并回到登录页。'),
('业务数据','读取失败时保留页面框架和重试入口；空数据与请求失败分别呈现，不用虚构指标替代真实响应。'),
('地图服务','无效坐标不绘制；SDK、Key 或域名配置问题显示地图状态提示，其他工单列表仍可访问。'),
('文件上传','格式、数量和大小错误在上传前反馈；网络失败保留队列；成功预览标记为临时状态，待后端同步后展示持久材料。'),
('Dify 嵌入','按需创建 iframe；加载超时提供重试和独立窗口入口。跨域环境不把 iframe load 事件当作模型服务健康证明。'),
('页面生命周期','切页时销毁控制器、监听器和定时器；刷新任务防止并发重入，减少重复请求和资源泄漏。')]
add_grid(d,['异常场景','处理方式'],rows,[1.55,4.59])
add_head(d,'七、实施路径与验证重点')
add_para(d,'实施时先确认后端字段、状态编码、认证方式和坐标系，再建立统一的数据映射；随后接入总览、事件、地图和设备模块；最后完成 COS 上传、Dify 嵌入和部署配置。接口字段变化优先在 api.js 边界适配，避免后端差异散落到页面组件。')
add_para(d,'验证重点包括：登录失败与会话过期处理；事件筛选与状态计数一致；审核或派单后列表更新；无坐标事件仍可从列表访问；地图聚合和图层筛选有效；上传限额、转码、进度和重试正常；Dify 加载失败时提供重试；页面离开后停止定时器。项目已有测试覆盖工单状态映射、地图聚合、设备状态、上传和缓存等逻辑。')
add_head(d,'八、前端模块索引')
rows=[
('页面入口','web/index.html、web/js/login.js'),
('应用控制器','web/js/app.js：导航、页面生命周期、总览刷新与状态协调'),
('数据访问','web/js/api.js、web/js/read-cache.js、web/js/types.js'),
('地图与设备','web/js/tencent-map.js、web/js/device-health.js、web/js/personnel-location.js'),
('工单与上传','web/js/order-center.js、web/js/finish-upload.js、web/js/cos-direct.js'),
('智能助手','web/js/patrol-agent.js、web/js/config.js'),
('构建与部署','scripts/build-web.mjs、server/web-finish-server.mjs、deploy/nginx')]
add_grid(d,['模块','主要文件与职责'],rows,[1.55,4.59])

# Apply FangSong GB2312, 小四 (12pt) consistently to body, table cells and captions.
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 try:
  st=d.styles[name]; st.font.name='仿宋_GB2312';st.font.size=Pt(12)
  st._element.rPr.rFonts.set(qn('w:eastAsia'),'仿宋_GB2312')
  st._element.rPr.rFonts.set(qn('w:ascii'),'仿宋_GB2312');st._element.rPr.rFonts.set(qn('w:hAnsi'),'仿宋_GB2312')
 except Exception:pass
for p in d.paragraphs:
 for r in p.runs:font(r,bold=r.bold)
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    for r in p.runs:font(r,bold=r.bold)
sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=Inches(1.18);sec.right_margin=Inches(1.18);sec.top_margin=Inches(.85);sec.bottom_margin=Inches(.85)
d.core_properties.title='智巡云 Web 端作品功能与原型设计 精进版'
d.core_properties.subject='作品功能、页面原型、业务流程与技术实现说明'
d.core_properties.author='智巡云项目组'
for table_index in [11,12]:
 row=d.tables[table_index].rows[0]
 trPr=row._tr.get_or_add_trPr()
 if not trPr.xpath('./w:tblHeader'): trPr.append(OxmlElement('w:tblHeader'))
# Add descriptive alternative text and keep each screenshot/figure with its caption.
alt=[
 '智巡云登录页面原型','路网总览首页截图','事件中心与工单详情截图','地图巡检告警点位截图',
 '设备中心截图','智巡助手浮窗截图','工单材料上传区域截图',
 'Web端系统架构图','告警审核与工单处置闭环图','数据刷新与缓存失效流程图','浏览器直传COS材料上传链路图'
]
for i,shape in enumerate(d.inline_shapes):
 if i < len(alt): shape._inline.docPr.set('descr',alt[i])
for t in d.tables:
 for row in t.rows:
  if row._tr.xpath('.//w:drawing'):
   if not row._tr.xpath('./w:trPr/w:cantSplit'):
    trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
d.save(out)
print(out)

