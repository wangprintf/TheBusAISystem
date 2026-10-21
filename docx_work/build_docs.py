
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import shutil, os

REF = Path(r'C:\Users\wangmeng\Desktop\智巡云-作品功能与原型设计(1).docx')
BASE = Path(r'C:\Users\wangmeng\Desktop\git\docx_work')
ASSET_ZY = BASE/'assets_zhixunyun'
ASSET_ZA = Path(r'C:\Users\wangmeng\Desktop\智眼路安-新工科比赛版\docx_assets')
OUT_ZY = Path(r'C:\Users\wangmeng\Desktop\智巡云-作品功能与原型设计 Web 版.docx')
OUT_ZA = Path(r'C:\Users\wangmeng\Desktop\智眼路安-作品功能与原型设计 Web 版.docx')

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', sz='6'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None:
            node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_font(run, name='微软雅黑', size=10.5, bold=None, color='000000'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def style_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, line=1.4):
    p.alignment = align
    pf=p.paragraph_format
    pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=line
    pf.widow_control=True

def add_body(doc, text, bold_label=None):
    p=doc.add_paragraph()
    style_paragraph(p)
    if bold_label and text.startswith(bold_label):
        r=p.add_run(bold_label); set_font(r,bold=True)
        r=p.add_run(text[len(bold_label):]); set_font(r)
    else:
        r=p.add_run(text); set_font(r)
    return p

def add_heading(doc, text, level=1):
    p=doc.add_paragraph()
    p.paragraph_format.keep_with_next=True
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=10 if level==1 else 7, after=8 if level==1 else 5, line=1.2)
    r=p.add_run(text); set_font(r,size=14 if level==1 else 12,bold=True)
    return p

def add_title(doc, text):
    p=doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=14, line=1.2)
    r=p.add_run(text); set_font(r,size=18,bold=True)
    return p

def add_meta(doc, label, value):
    p=doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=3, line=1.35)
    r=p.add_run(label); set_font(r,bold=True)
    r=p.add_run(value); set_font(r)
    return p

def add_feature_table(doc, rows):
    table=doc.add_table(rows=1, cols=2)
    table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    widths=[Inches(1.55),Inches(4.62)]
    for i,w in enumerate(widths):
        table.columns[i].width=w
    hdr=table.rows[0].cells
    for i,t in enumerate(['功能名称','功能描述']):
        hdr[i].text=''
        p=hdr[i].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(t); set_font(r,bold=True)
        set_cell_shading(hdr[i],'DCE6F1'); set_cell_border(hdr[i]); set_cell_margins(hdr[i])
        hdr[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for name,desc in rows:
        cells=table.add_row().cells
        for i,t in enumerate([name,desc]):
            cells[i].text=''
            p=cells[i].paragraphs[0]
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.15
            r=p.add_run(t); set_font(r,size=9.5,bold=(i==0))
            set_cell_border(cells[i]); set_cell_margins(cells[i],top=90,bottom=90)
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return table

def remove_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            tcPr=cell._tc.get_or_add_tcPr()
            borders=tcPr.first_child_found_in('w:tcBorders')
            if borders is None:
                borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
            for edge in ('top','left','bottom','right','insideH','insideV'):
                el=borders.find(qn('w:'+edge))
                if el is None:
                    el=OxmlElement('w:'+edge); borders.append(el)
                el.set(qn('w:val'),'nil')

def add_figure(doc, image_path, caption, width=6.0):
    table=doc.add_table(rows=1, cols=1)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    cell=table.cell(0,0); set_cell_margins(cell,top=60,bottom=60,start=60,end=60)
    remove_table_borders(table)
    p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
    run=p.add_run(); run.add_picture(str(image_path), width=Inches(width))
    p2=cell.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.space_after=Pt(4)
    r=p2.add_run(caption); set_font(r,size=9.5,bold=True)
    return table

def add_two_figures(doc, left_path, left_caption, right_path, right_caption, width=3.05):
    table=doc.add_table(rows=1, cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    table.columns[0].width=Inches(width); table.columns[1].width=Inches(width)
    for cell,path,cap in [(table.cell(0,0),left_path,left_caption),(table.cell(0,1),right_path,right_caption)]:
        set_cell_margins(cell,top=60,bottom=60,start=60,end=60)
        p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path),width=Inches(width-0.1))
        p2=cell.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p2.add_run(cap); set_font(r,size=9,bold=True)
    remove_table_borders(table)
    return table

def add_page_break(doc):
    p=doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

def init_from_template():
    doc=Document(str(REF))
    body=doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    sec=doc.sections[0]
    sec.top_margin=Inches(1.0); sec.bottom_margin=Inches(1.0); sec.left_margin=Inches(1.18); sec.right_margin=Inches(1.18)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='微软雅黑'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑'); normal.font.size=Pt(10.5)
    return doc

def build_zhixunyun():
    doc=init_from_template()
    add_title(doc,'智巡云 Web 端作品功能与原型设计')
    add_body(doc,'“智巡云——基于边缘感知的路网健康巡检系统”Web 管理后台面向交通管理部门的路网巡检、告警审核、工单处置与设备运维场景，围绕路网总览、事件中心、地图巡检、设备中心、材料上传和智能助手构建一体化工作台。下面从功能概述与原型设计两个方面说明网页端的功能边界、交互流程和界面原型。')
    add_heading(doc,'功能概述',1)
    rows=[
      ('登录认证','用户通过账号和密码登录，前端调用 /api/login 获取令牌并保存会话资料；会话失效或退出登录时清除本地状态并返回登录页。'),
      ('路网总览','首页集中展示今日告警、待处理工单、设备在线率、处置趋势、设备状态和地图告警，支持定时刷新和关键指标原位更新。'),
      ('事件中心','以队列和详情双栏呈现事件，支持关键词、状态、优先级、日期和排序筛选，提供审核、发送工单、催办和回查。'),
      ('地图巡检','接入腾讯地图 GL JS，按坐标绘制告警点、聚合点、巡检路线和人员位置，点击标记可定位事件详情。'),
      ('设备中心','按正常、异常、离线状态统计车载设备，展示设备编号、车牌、温度、存储和异常原因，并支持明细弹窗。'),
      ('处理材料','工单处理中可多选图片和视频，浏览器校验格式与大小，PNG/WEBP 转 JPG 后直传腾讯云 COS，逐项展示进度并支持重试。'),
      ('数据同步','api.js 统一处理 REST 请求、字段映射、短缓存和写后失效；页面切换和异步刷新使用导航编号避免旧响应覆盖新数据。'),
      ('智巡助手','通过 Dify iframe 延迟加载 PatrolCloud 智巡助手，支持展开、拖拽、调整尺寸、重试、关闭和跨域会话保留。'),
      ('个人信息','头像入口展示登录接口返回的姓名与部门，支持退出登录；个人信息随会话变化自动更新。'),
      ('发布与部署','网页可通过 npx serve 或 Node 静态服务运行，构建脚本输出 web-dist，生产环境使用 Nginx 发布并配置地图、COS 和 Dify。'),
    ]
    add_feature_table(doc,rows)
    add_heading(doc,'原型设计',1)
    add_meta(doc,'软件名称：','智巡云——基于边缘感知的路网健康巡检系统')
    add_meta(doc,'软件类型：','智慧城市 / 交通巡检管理 Web 端')
    add_meta(doc,'实现平台：','桌面浏览器为主，兼容平板与移动端响应式访问')
    add_meta(doc,'技术架构：','原生 HTML5、CSS3、JavaScript ES Modules；通过 REST API 对接业务服务，浏览器直传腾讯云 COS，Dify 以 iframe 接入')
    add_meta(doc,'地图与云服务：','腾讯位置服务 GL JS、腾讯云 COS 对象存储、Dify 智能体')
    add_meta(doc,'本地运行：','npx serve web 或 npm run start:web；默认 Node 预览端口为 8290')
    add_meta(doc,'构建方式：','npm run build:web，生成 web-dist 静态部署目录')
    add_body(doc,'系统原型按业务流程依次包括登录页、路网总览、事件中心、地图巡检、设备中心、工单材料上传和智巡助手等界面。以下说明结合当前 Web 页面和实际功能实现展开。')
    add_body(doc,'登录页提供账号、密码输入和错误提示，登录成功后才动态加载业务应用。页面采用深色科技风格，品牌区与凭证表单分区展示，便于用户快速识别系统入口。')
    add_figure(doc,ASSET_ZY/'01-login.png','图 1 登录界面',6.0)
    add_body(doc,'路网总览以驾驶舱方式组织数据，左侧展示工单状态，中部呈现路网地图，右侧显示运行概况、设备状态和趋势分析。总览页并行加载多类数据，并通过定时刷新更新指标和图表。')
    add_figure(doc,ASSET_ZY/'02-dashboard.png','图 2 路网总览首页',6.0)
    add_body(doc,'事件中心采用“报警处置队列 + 工单详情”双栏结构。左侧可按状态、优先级、日期和关键词筛选，右侧展示事件材料、处理信息和操作按钮；高优先级事件使用醒目颜色提示。')
    add_figure(doc,ASSET_ZY/'03-orders.png','图 3 事件中心与工单详情',6.0)
    add_body(doc,'地图巡检页以腾讯地图为核心，红色和黄色点位区分优先级，邻近点位自动聚合，点击单点可回到事件详情。地图同时提供图层筛选和人员位置显示，支持从空间维度研判事件分布。')
    add_figure(doc,ASSET_ZY/'04-map.png','图 4 地图巡检与告警点位',4.9)
    add_body(doc,'设备中心汇总设备总数及正常、异常、离线状态，设备列表展示车载设备编号、车辆名称、状态、温度和存储等信息。点击状态分布或列表项可查看设备明细和异常原因。')
    add_figure(doc,ASSET_ZY/'05-devices.png','图 5 设备中心',6.0)
    add_body(doc,'工单处理材料区域支持多选图片和视频，逐项反馈上传进度。浏览器会校验文件类型、数量和大小，并把 PNG/WEBP 转成 JPG 后直接上传到 COS；失败项保留在队列中，可再次重试。智巡助手从页面入口按需展开，避免首屏等待。')
    add_two_figures(doc,ASSET_ZY/'06-assistant.png','图 6 智巡助手浮窗',ASSET_ZY/'03-orders.png','图 7 材料上传所在工单详情',3.05)
    add_body(doc,'个人信息入口位于页面右上角，显示姓名、部门和退出登录按钮；所有用户资料来自登录会话。页面切换时会销毁当前控制器和定时器，避免后台页面继续刷新。')
    return doc

def build_zhiyan():
    doc=init_from_template()
    add_title(doc,'智眼路安 Web 端作品功能与原型设计')
    add_body(doc,'“智眼路安——基于视觉感知与边缘计算的路网健康管理平台”是面向新工科实践创新大赛的 Web 管理后台。平台围绕视觉感知、边缘计算和协同管养，将路网告警、工单处置、地图态势、设备运维、处理材料和路安助手整合到统一驾驶舱。本文沿用参考文档的“功能概述 + 原型设计”结构，说明比赛版的页面功能、交互链路和主题化界面。')
    add_heading(doc,'功能概述',1)
    rows=[
      ('比赛版驾驶舱','使用深森林绿、薄荷青和暖金配色，桌面端采用左侧工单、中间地图、右侧分析的三栏布局，移动端自动纵向排列。'),
      ('登录与会话','登录页突出“视觉感知、边缘计算、协同管养”能力文案，调用 /api/login 验证账号，并将令牌、姓名和部门保存到 sessionStorage。'),
      ('路网总览','聚合告警、工单和设备指标，展示工单状态、地图态势、设备健康和近 7 日趋势，支持定时刷新。'),
      ('事件中心','支持关键词、日期、优先级、处置状态筛选，详情页提供审核、发送工单、催办、材料上传和结果回查。'),
      ('地图巡检','腾讯地图 GL JS 绘制告警点和聚合点，独立维护人员图层，支持点位点击、图层筛选和空间态势浏览。'),
      ('设备中心','按正常、异常、离线分布展示车载设备，支持状态筛选、设备详情、异常原因和运行指标查看。'),
      ('处理材料','图片、视频在浏览器完成格式和大小校验，PNG/WEBP 转 JPG 后直传 COS，逐项反馈进度，失败项可重试。'),
      ('路安助手','右侧“路安助手”通过 Dify iframe 延迟加载，支持拖动、缩放、关闭、重试、独立窗口和收起后保留会话。'),
      ('接口与缓存','复用稳定的 api.js、read-cache.js、字段映射和写后失效机制，保证比赛版与原有后端数据链路一致。'),
      ('运行与发布','npm run start:web 默认以 8291 端口启动比赛版，npm run build:web 生成独立 web-dist 静态部署产物。'),
    ]
    add_feature_table(doc,rows)
    add_heading(doc,'原型设计',1)
    add_meta(doc,'软件名称：','智眼路安——基于视觉感知与边缘计算的路网健康管理平台')
    add_meta(doc,'软件类型：','智慧城市 / 视觉感知与边缘计算路网管理 Web 端')
    add_meta(doc,'实现平台：','桌面浏览器、平板和移动端响应式访问；比赛版本地端口为 8291')
    add_meta(doc,'技术架构：','原生 HTML5、CSS3、JavaScript ES Modules；业务逻辑复用原项目，表现层由 competition-theme.css 统一改造')
    add_meta(doc,'地图与云服务：','腾讯位置服务 GL JS、腾讯云 COS 对象存储、Dify 路安助手')
    add_meta(doc,'数据关系：','比赛版默认连接与原项目相同的 API、数据库、地图、COS 和 Dify 服务；写操作会影响同一套业务数据')
    add_meta(doc,'构建方式：','npm run build:web，生成比赛版 web-dist 静态部署目录')
    add_body(doc,'系统原型按“登录—总览—事件—地图—设备—材料—智能助手”的工作流组织。比赛版主要重构页面主题、导航位置和信息布局，接口、地图控制、工单流程、设备操作与上传逻辑沿用稳定实现。')
    add_body(doc,'登录界面将凭证表单放在左侧，将产品名称、视觉感知与边缘计算说明放在右侧，使用森林绿背景、暖金分隔线和浅色表单卡片形成比赛版入口。')
    add_figure(doc,ASSET_ZA/'01-login.png','图 1 智眼路安登录界面',6.0)
    add_body(doc,'登录后驾驶舱采用三栏信息架构：左侧为工单状态和处置队列，中部为城市路网态势地图，右侧为运行概况、设备状态和趋势分析。该布局适合大屏讲解，也通过 CSS Grid 断点适配较窄屏幕。')
    add_figure(doc,ASSET_ZA/'02-dashboard-assistant.png','图 2 比赛版驾驶舱与路安助手',6.0)
    add_body(doc,'路网总览通过状态环图、地图标记和运行指标建立统一的空间与业务视图。下图为依据当前页面结构整理的界面原型示意，用于说明工单、地图和分析模块的相对位置与信息层级。')
    add_figure(doc,ASSET_ZA/'03-dashboard-wireframe.png','图 3 路网总览界面原型示意',6.0)
    add_body(doc,'事件中心把事件队列和详情面板放在同一工作区，优先展示未处理和高等级工单；详情侧集中呈现事件类型、位置、报警材料、处理后材料以及发送、审核和上传操作。')
    add_figure(doc,ASSET_ZA/'04-orders-wireframe.png','图 4 事件中心界面原型示意',6.0)
    add_body(doc,'设备中心通过状态分布和设备列表呈现车载设备健康度。正常、异常和离线设备采用不同状态色，点击状态区域可查看设备清单与异常原因，适合在比赛现场展示设备运维闭环。')
    add_figure(doc,ASSET_ZA/'05-devices-wireframe.png','图 5 设备中心界面原型示意',6.0)
    add_body(doc,'处理材料页把文件选择、逐项进度和失败重试集中在一个面板中。上传前执行格式、数量和大小校验，成功后先显示临时预览，原有扫描与查询接口随后同步持久材料。')
    add_figure(doc,ASSET_ZA/'06-upload-wireframe.png','图 6 工单材料上传界面原型示意',6.0)
    add_body(doc,'路安助手从页面右上角入口打开，以 Dify iframe 承载对话。组件支持拖动、调整尺寸、Esc 关闭和重试；跨域场景下只依据加载、超时和错误事件提示连接状态，不虚构模型服务健康状态。')
    add_body(doc,'比赛版与原项目共用后端数据，因此演示环境应提前确认账号、API、地图 Key、COS CORS 和 Dify 嵌入地址。若需要与原版完全隔离，应另行配置数据库、存储桶和服务地址。')
    return doc

for builder,out in [(build_zhixunyun,OUT_ZY),(build_zhiyan,OUT_ZA)]:
    doc=builder()
    doc.core_properties.title=out.stem
    doc.core_properties.subject='Web 端作品功能与原型设计'
    doc.core_properties.author='智巡云项目组'
    doc.save(str(out))
    print(out)

