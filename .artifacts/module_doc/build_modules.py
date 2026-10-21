from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
ROOT.mkdir(parents=True, exist_ok=True)
OUTPUT = Path(r'C:\Users\wangmeng\Desktop\智巡云_Web端核心模块与实现说明.docx')
FONT = r'C:\Windows\Fonts\simfang.ttf'
INK, LINE, BLUE, PALE = '#24282D', '#727B84', '#405F79', '#F2F5F7'

def font(n):
    return ImageFont.truetype(FONT, n)

def canvas():
    im = Image.new('RGB', (1600, 730), 'white')
    return im, ImageDraw.Draw(im)

def text(d, xy, s, size=36, fill=INK, anchor='mm'):
    d.text(xy, s, font=font(size), fill=fill, anchor=anchor)

def block(d, x, y, w, h, title, lines, number=None):
    d.rectangle((x, y, x+w, y+h), outline=LINE, width=2, fill='white')
    if number:
        text(d, (x+26, y+25), number, 27, BLUE, 'lt')
    text(d, (x+w/2, y+62), title, 43, BLUE)
    for i, line in enumerate(lines):
        text(d, (x+w/2, y+126+i*49), line, 35)

def arrow(d, points):
    d.line(points, fill=LINE, width=3)
    x,y = points[-1]; px,py = points[-2]
    if x>px: head=[(x,y),(x-14,y-8),(x-14,y+8)]
    elif x<px: head=[(x,y),(x+14,y-8),(x+14,y+8)]
    elif y>py: head=[(x,y),(x-8,y-14),(x+8,y-14)]
    else: head=[(x,y),(x-8,y+14),(x+8,y+14)]
    d.polygon(head, fill=LINE)

def strip(d, y, label, values):
    text(d, (45,y+35), label, 36, BLUE, 'lm')
    for i,v in enumerate(values):
        x=300+i*425
        d.rectangle((x,y,x+390,y+70), outline='#BAC2C9', width=2, fill=PALE)
        text(d,(x+195,y+35),v,34)

im,d=canvas()
block(d,45,40,440,275,'页面模块',['总览 · 事件','地图 · 设备','app.js 组织导航与交互'])
block(d,580,40,440,275,'数据访问',['api.js 统一请求与映射','read-cache.js 短时复用'])
block(d,1115,40,440,275,'业务服务',['REST API','告警 · 工单 · 设备数据'])
arrow(d,[(485,176),(580,176)]); arrow(d,[(1020,176),(1115,176)])
text(d,(800,365),'页面呈现依赖统一数据访问，业务状态以服务端返回为准',34)
strip(d,430,'刷新控制',['短缓存与请求复用','写入成功后重新读取','切页校验与停止刷新'])
strip(d,570,'外部接入',['腾讯地图  点位展示','COS  材料存储','Dify  问答入口'])
im.crop((0,0,1600,665)).save(ROOT/'module_1.png',dpi=(240,240))

im,d=canvas()
xs=[45,440,835,1230]
steps=[('事件复核',['查看告警与证据','人工判定有效性']),('发送工单',['确认处置任务','状态转为处理中']),('现场处置',['填写处理说明','上传图片或视频']),('结果复核',['待审核时检查材料','合格后标记已完成'])]
for i,(title,lines) in enumerate(steps):
    block(d,xs[i],95,325,265,title,lines,f'0{i+1}')
for i in range(3): arrow(d,[(xs[i]+325,225),(xs[i+1],225)])
text(d,(405,195),'有效',28)
text(d,(207,405),'无效事件记录审核结果',29)
arrow(d,[(1392,360),(1392,530),(602,530),(602,360)])
d.rectangle((730,500,1265,560),fill='white')
text(d,(997,530),'待审核时复核不合格，返回未处理',33)
text(d,(800,648),'每次操作由接口确认，再更新处置队列和工单详情',35)
im.save(ROOT/'module_2.png',dpi=(240,240))

im,d=canvas()
xs=[45,580,1115]
items=[('选择与校验',['处理中的工单可上传','检查格式、数量和大小']),('格式处理',['PNG / WEBP 转白底 JPG','JPG 与视频按规则处理']),('直传 COS',['逐项展示上传进度','失败项保留并可重试'])]
for i,(title,lines) in enumerate(items): block(d,xs[i],65,440,265,title,lines,f'0{i+1}')
arrow(d,[(485,198),(580,198)]); arrow(d,[(1020,198),(1115,198)])
text(d,(800,389),'每批最多 10 个文件，单文件不超过 200 MB',35)
d.line([(45,442),(1555,442)],fill='#D7DCE0',width=2)
for i,(title,lines) in enumerate([('临时预览',['上传成功后立即显示']),('扫描与查询',['既有流程同步材料记录']),('工单展示',['从查询结果加载持久材料'])]):
    block(d,xs[i],490,440,180,title,lines)
arrow(d,[(485,585),(580,585)]); arrow(d,[(1020,585),(1115,585)])
im.save(ROOT/'module_3.png',dpi=(240,240))

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.82);sec.right_margin=Inches(.82)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)

def set_font(obj):
    obj.font.name='仿宋_GB2312';obj.font.size=Pt(12);obj.font.color.rgb=RGBColor.from_string('000000')
    rp=obj.element.get_or_add_rPr() if hasattr(obj,'element') else obj._element.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for k in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+k),'仿宋_GB2312')
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: rf.attrib.pop(qn('w:'+k),None)

for name in ['Normal','Title','Heading 1','Heading 2','Caption','Header','Footer']:
    st=doc.styles[name];set_font(st)
    pf=st.paragraph_format;pf.space_after=Pt(4);pf.space_before=Pt(0);pf.line_spacing=1.3
    if name.startswith('Heading'): st.font.bold=True;pf.space_before=Pt(9);pf.keep_with_next=True
doc.styles['Normal'].paragraph_format.first_line_indent=Pt(24)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Caption'].paragraph_format.first_line_indent=Pt(0)
doc.styles['Caption'].font.bold=False
for st in doc.styles:
    for borders in list(st.element.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)

def para(s,style=None):
    p=doc.add_paragraph(s,style)
    for r in p.runs:set_font(r)
    p.paragraph_format.widow_control=True
    return p

def heading(s): return para(s,'Heading 1')
def sub(s): return para(s,'Heading 2')
def picture(n,caption):
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
    p.add_run().add_picture(str(ROOT/f'module_{n}.png'),width=Inches(6.86))
    p=para(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(8)

p=para('智巡云 Web 端核心模块与实现说明','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
para('智巡云 Web 端面向路网巡检管理，提供态势查看、事件审核、工单处置和设备查询。本说明将原有四张图整合为三个模块，介绍功能关系与实现方式，可用于作品功能与原型设计。')
heading('一 系统架构与数据刷新模块')
para('Web 端采用 HTML5、CSS3 和原生 JavaScript ES Modules 实现。app.js 负责页面导航和交互组织，api.js 集中完成业务请求、认证和字段映射，read-cache.js 复用短时间内的读取结果。总览、事件、地图与设备页面共用这一数据入口，并分别接入腾讯地图、COS 和 Dify。操作成功后重新读取相关数据，切换页面时校验响应归属，使展示结果与当前任务保持一致。')
picture(1,'图 1 页面模块 数据访问与刷新控制')
sub('页面职责与外部服务')
para('路网总览汇总运行情况；事件中心集中展示处置队列和工单详情；地图页呈现有效坐标与人员位置；设备页展示状态及异常说明。腾讯地图承担点位绘制，COS 保存处理材料，Dify 通过按需创建的 iframe 提供问答入口。')
sub('刷新与异常处理')
para('读取缓存以 API 地址和登录会话为作用域，默认有效期为 1 秒，并复用进行中的同类请求。写入成功后使相关缓存失效，再加载最新数据。异步响应返回时检查导航编号和数据版本；离开页面后停止对应刷新任务。请求失败时显示重试入口，地图无效坐标不绘制，相关事件仍可从列表查看。')

doc.add_page_break()
heading('二 事件审核与工单处置模块')
para('事件中心以处置队列和工单详情组成工作区。管理人员先查看告警地点、上报时间和证据，完成人工复核，再发送工单并跟进现场处置。处理说明及图片、视频集中保存在详情中，供结果审核时回查。待审核工单通过复核后转为已完成，不合格则返回未处理。页面操作由业务接口确认，成功后刷新列表和详情；失败时保留当前选择和筛选条件，方便继续处理。')
picture(2,'图 2 事件复核到工单完成的处理流程')
sub('队列与详情保持同步')
para('队列支持关键词、状态、优先级、日期和排序筛选。点击记录后，详情区展示对应事件、材料和处理进度；状态计数、可见列表与选中项同步更新。空列表、无筛选结果和接口失败分别提示，状态与优先级采用文字配合颜色表达。')
sub('状态写入与结果回查')
para('发送工单调用事件更新接口，将状态写为处理中，并追加操作备注；催办追加催办记录。待审核工单的复核操作写入审核结果，通过时设为已完成，不通过时设为未处理。负责人信息从分配数据读取，实际分配结果以接口返回为准。现场处置与提交审核可由其他业务端完成，Web 端负责查看、操作和回查相应状态。')
sub('模块衔接')
para('工单操作完成后重新读取相关资源，使队列、详情与总览统计更新。材料上传入口仅在处理中开放；上传进度和失败信息在当前详情内反馈，避免管理人员反复切换页面确认结果。')

doc.add_page_break()
heading('三 工单处理材料上传模块')
para('处理材料上传用于补充现场处置凭证。用户在处理中的工单内选择图片或视频，浏览器先检查文件格式、数量和大小，再将 PNG、WEBP 转换为白底 JPG，随后直接上传至腾讯云 COS。队列逐项显示进度，失败文件保留以便重试。上传成功后先提供临时预览，既有扫描流程和查询接口再完成材料记录同步，工单详情据此展示可回查的处理材料。')
picture(3,'图 3 文件上传与工单材料同步')
sub('上传前检查与格式处理')
para('常用格式为 JPG、PNG、WEBP 图片和 MP4 视频，每批最多 10 个文件，单文件不超过 200 MB；空文件及不匹配的文件类型在上传前反馈。PNG、WEBP 使用浏览器图像能力转换，透明区域填充白色，转换后再次检查大小；过大分辨率或无法读取的图片提示用户缩小或改用 JPG。')
sub('传输进度与失败重试')
para('文件通过 XMLHttpRequest 直传 COS，对象键包含工单编号和 UUID，以区分材料并支持现有扫描流程关联工单。上传过程中逐项反馈进度，暂时禁用队列移除和重复提交；失败时显示具体原因，重试仅处理未成功项。用户需要保留当前窗口，等待本批上传结束。')
sub('预览与持久材料的关系')
para('传输成功表示文件已进入对象存储。页面立即补充临时预览，材料记录的持久同步仍依赖既有扫描和查询流程。再次获取工单时，将查询结果与临时预览按文件地址合并去重，避免同一材料重复展示；工单是否完成仍由审核和业务状态决定。')

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Pt(0)
r=footer.add_run();set_font(r)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='智巡云 Web 端核心模块与实现说明'
doc.core_properties.subject='系统架构与数据刷新 事件工单处置 工单材料上传'
doc.core_properties.author=''
# The requested legacy font name is retained; the installed FangSong is an explicit fallback.
for part in doc.part.package.parts:
    if str(part.partname)=='/word/fontTable.xml':
        from lxml import etree
        root=etree.fromstring(part.blob)
        f=OxmlElement('w:font');f.set(qn('w:name'),'仿宋_GB2312')
        alt=OxmlElement('w:altName');alt.set(qn('w:val'),'仿宋');f.append(alt);root.append(f)
        part._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
doc.save(OUTPUT)
print(OUTPUT)
